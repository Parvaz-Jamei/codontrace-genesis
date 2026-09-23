"""English BAIC manuscript in the congress Word template.

The page, styles, and numbering come from the official congress file
(paper/baic/baic_congress_template.docx), not from an IEEE template.
English paragraphs are left-to-right. The saved package has one
core-properties part, which is what Word requires.
"""

from __future__ import annotations

import io
import zipfile
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.shared import Inches
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "paper" / "baic" / "baic_congress_template.docx"
OUT = ROOT / "paper" / "baic" / "BAIC2026_Jamei_en.docx"
FIG = ROOT / "paper" / "baic" / "figures" / "audit_flow.png"
COL = 4600
PKG = "http://schemas.openxmlformats.org/package/2006/relationships"


def _flow_png(path: Path) -> None:
    """Column-width figure. Type is large enough to read at 3.05 inches."""
    try:
        from PIL import Image, ImageDraw, ImageFont
    except ImportError:
        if path.is_file():
            return
        raise
    serif = "/usr/share/fonts/truetype/liberation/LiberationSerif-Regular.ttf"
    bold = "/usr/share/fonts/truetype/liberation/LiberationSerif-Bold.ttf"
    title_font = ImageFont.truetype(bold, 44)
    sub_font = ImageFont.truetype(serif, 32)
    label_font = ImageFont.truetype(bold, 36)
    width, height = 980, 980
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    steps = (
        ("Domain profile", "life, biomedical, or hardware"),
        ("Evidence bundle", "arms, replay, labels"),
        ("Grade 0 to 5", "the bundle, not the engine"),
        ("Risk bar", "does not move the grade"),
    )
    box_w, box_h, gap = 920, 118, 36
    x = (width - box_w) // 2
    y = 8
    for index, (title, sub) in enumerate(steps):
        draw.rounded_rectangle((x, y, x + box_w, y + box_h), radius=12, outline="black", width=4)
        if index == 3:
            draw.rounded_rectangle((x + 8, y + 8, x + box_w - 8, y + box_h - 8), radius=8, outline="black", width=2)
        tw = draw.textlength(title, font=title_font)
        sw = draw.textlength(sub, font=sub_font)
        draw.text(((width - tw) / 2, y + 16), title, fill="black", font=title_font)
        draw.text(((width - sw) / 2, y + 68), sub, fill="black", font=sub_font)
        if index < len(steps) - 1:
            ax = width // 2
            y2 = y + box_h
            tip = y2 + gap - 2
            if index == 2:
                yy = y2 + 3
                while yy < tip - 12:
                    draw.line((ax, yy, ax, min(yy + 8, tip - 12)), fill="black", width=4)
                    yy += 14
            else:
                draw.line((ax, y2 + 2, ax, tip - 12), fill="black", width=4)
            draw.polygon([(ax - 10, tip - 14), (ax + 10, tip - 14), (ax, tip)], fill="black")
        y += box_h + gap
    top = y + 10
    draw.line((x, top, x + box_w, top), fill="black", width=3)
    head = "Recorded grades"
    draw.text(((width - draw.textlength(head, font=title_font)) / 2, top + 12), head, fill="black", font=title_font)
    bar_x, bar_max = 340, 480
    base = top + 78
    for name, grade in (("HE01", 4), ("Device table", 0)):
        draw.text((x + 16, base), name, fill="black", font=label_font)
        draw.rectangle((bar_x, base + 4, bar_x + bar_max, base + 40), outline="black", width=3)
        fill_w = int(bar_max * grade / 5)
        if fill_w:
            draw.rectangle((bar_x + 3, base + 7, bar_x + fill_w - 3, base + 37), fill="black")
        draw.text((bar_x + bar_max + 14, base), str(grade), fill="black", font=label_font)
        base += 58
    path.parent.mkdir(parents=True, exist_ok=True)
    mask = Image.eval(image.convert("L"), lambda pixel: 255 - pixel)
    box = mask.getbbox()
    if box:
        pad = 8
        image = image.crop((
            max(0, box[0] - pad),
            max(0, box[1] - pad),
            min(width, box[2] + pad),
            min(height, box[3] + pad),
        ))
    image.save(path, "PNG", dpi=(300, 300))


def _figure(doc: Document, path: Path) -> None:
    paragraph = doc.add_paragraph()
    _ltr(paragraph, "center")
    paragraph.paragraph_format.space_before = 0
    paragraph.paragraph_format.space_after = 0
    run = paragraph.add_run()
    run.add_picture(str(path), width=Inches(3.05))
    doc_pr = run._r.find(".//" + qn("wp:docPr"))
    if doc_pr is not None:
        doc_pr.set("name", "Figure 1")
        doc_pr.set(
            "descr",
            "Flow from domain profile to grade, then the two recorded grades: HE01 is 4 and the device table is 0.",
        )


def _styles(doc: Document) -> dict:
    return {style.style_id: style for style in doc.styles}


def _ltr(paragraph, align: str = "both") -> None:
    ppr = paragraph._p.get_or_add_pPr()
    bidi = ppr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        ppr.append(bidi)
    bidi.set(qn("w:val"), "0")
    jc = ppr.find(qn("w:jc"))
    if jc is None:
        jc = OxmlElement("w:jc")
        ppr.append(jc)
    jc.set(qn("w:val"), align)


def _times(run, size_pt: float, bold: bool = False, italic: bool = False) -> None:
    run.bold = bold
    run.italic = italic
    run.font.name = "Times New Roman"
    rpr = run._r.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), "Times New Roman")
    half = str(int(size_pt * 2))
    for tag in ("w:sz", "w:szCs"):
        node = rpr.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            rpr.append(node)
        node.set(qn("w:val"), half)
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(qn("w:val"), "en-US")
    lang.set(qn("w:bidi"), "en-US")


def _p(doc: Document, styles: dict, style_id: str, text: str = "", size: float = 10, align: str = "both", bold: bool = False):
    paragraph = doc.add_paragraph()
    paragraph.style = styles[style_id]
    _ltr(paragraph, align)
    if text:
        run = paragraph.add_run(text)
        _times(run, size, bold=bold)
    if style_id == "Caption":
        paragraph.paragraph_format.keep_with_next = True
    return paragraph


def _drop_section_rtl(sect) -> None:
    for child in list(sect):
        if child.tag in {qn("w:bidi"), qn("w:rtlGutter")}:
            sect.remove(child)


def _clear(doc: Document):
    body = doc.element.body
    title_sect = None
    for child in list(body):
        ppr = child.find(qn("w:pPr")) if child.tag == qn("w:p") else None
        if ppr is not None:
            sect = ppr.find(qn("w:sectPr"))
            if sect is not None and title_sect is None:
                title_sect = deepcopy(sect)
    final = body.find(qn("w:sectPr"))
    if title_sect is None or final is None:
        raise SystemExit("congress template is missing its section break")
    _drop_section_rtl(title_sect)
    _drop_section_rtl(final)
    for child in list(body):
        if child is not final:
            body.remove(child)
    return title_sect


def _borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "000000")
        borders.append(element)
    tbl_pr.append(borders)
    width = OxmlElement("w:tblW")
    width.set(qn("w:w"), str(COL))
    width.set(qn("w:type"), "dxa")
    tbl_pr.append(width)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)


def _cell(cell, text: str, header: bool) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    _ltr(paragraph, "center" if header else "both")
    ppr = paragraph._p.get_or_add_pPr()
    ind = ppr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        ppr.append(ind)
    ind.set(qn("w:left"), "0")
    ind.set(qn("w:right"), "0")
    ind.set(qn("w:firstLine"), "0")
    ind.set(qn("w:hanging"), "0")
    run = paragraph.add_run(text)
    _times(run, 8, bold=header)


def _table(doc: Document, rows: list[list[str]], widths: list[int]) -> None:
    table = doc.add_table(rows=len(rows), cols=len(widths))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _borders(table)
    grid = table._tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for col, width in zip(grid.findall(qn("w:gridCol")), widths):
            col.set(qn("w:w"), str(width))
    for r_index, row in enumerate(rows):
        for c_index, text in enumerate(row):
            _cell(table.rows[r_index].cells[c_index], text, r_index == 0)
            tc_pr = table.rows[r_index].cells[c_index]._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[c_index]))
            tc_w.set(qn("w:type"), "dxa")


def _header(doc: Document) -> None:
    line = (
        "1st National Congress on Biomedical Engineering, "
        "Artificial Intelligence and Cognitive Science, "
        "Ferdowsi University of Mashhad, 2026"
    )
    for section in doc.sections:
        header = section.header
        header.is_linked_to_previous = False
        if not header.paragraphs:
            continue
        first = header.paragraphs[0]
        for paragraph in header.paragraphs:
            for run in paragraph.runs:
                run.text = ""
        run = first.add_run(line) if not first.runs else first.runs[0]
        run.text = line
        _times(run, 8)
        _ltr(first, "center")


def _sanitize(path: Path) -> None:
    """Drop the duplicate core.xml that makes Word call the file corrupt."""

    source = zipfile.ZipFile(path)
    buffer = io.BytesIO()
    seen: set[str] = set()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as out:
        for info in source.infolist():
            if info.filename in seen:
                continue
            data = source.read(info.filename)
            if info.filename == "_rels/.rels":
                root = etree.fromstring(data)
                kept = False
                for rel in list(root):
                    if not rel.tag.endswith("Relationship"):
                        continue
                    if (rel.get("Type") or "").endswith("/core-properties"):
                        if kept:
                            root.remove(rel)
                        kept = True
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            seen.add(info.filename)
            out.writestr(info, data)
    path.write_bytes(buffer.getvalue())
    check = zipfile.ZipFile(path)
    names = check.namelist()
    if len(names) != len(set(names)):
        raise SystemExit("duplicate zip parts remain")
    rels = etree.fromstring(check.read("_rels/.rels"))
    cores = [
        rel for rel in rels
        if rel.tag.endswith("Relationship") and (rel.get("Type") or "").endswith("/core-properties")
    ]
    if len(cores) != 1:
        raise SystemExit(f"expected one core-properties link, found {len(cores)}")


def build() -> None:
    if not TEMPLATE.is_file():
        raise SystemExit(f"missing congress template: {TEMPLATE}")
    doc = Document(str(TEMPLATE))
    styles = _styles(doc)
    title_sect = _clear(doc)
    _header(doc)

    _p(doc, styles, "Title", "A File Label Is Not an Audit Level", 14, "center", True)
    author = _p(doc, styles, "Author", "", 11, "center")
    name = author.add_run("Parvaz Jamei")
    _times(name, 11, bold=True)
    mark = author.add_run("1")
    _times(mark, 11)
    mark.font.superscript = True
    aff = _p(doc, styles, "Author", "", 9, "center")
    one = aff.add_run("1")
    _times(one, 9)
    one.font.superscript = True
    rest = aff.add_run(
        "Independent researcher, Mashhad, Iran, parvaz.jamie@gmail.com, "
        "ORCID 0009-0002-9980-270X"
    )
    _times(rest, 9)

    abstract = (
        "Abstract- Credibility guides say what evidence a computational model "
        "should carry. They do not grade one file from its declared label. "
        "This paper grades the bundle. The label intervention_supported "
        "maps only to an internal grade of 3. Public grade 4 is recorded only "
        "when replay is verified and an interval is present on at least 16 "
        "seeds. Thirty paired seeds reach grade 4. The contrast of 16.49 is "
        "receiver energy in the simulation, from two settings changed together, "
        "not a clinical result. A device score table that carries "
        "regulatory labels and no recorded run stays at grade 0. A declared "
        "model risk has a separate bar and does not move the grade. Six "
        "deletions of a ladder condition fall below grade 4. That count is not "
        "a rate over the published literature."
    )
    _p(doc, styles, "abstract0", abstract, 9, "both")
    keywords = _p(
        doc,
        styles,
        "Index",
        "Keywords- evidence audit, computational model credibility, unit of analysis, credibility ladder",
        9,
        "both",
    )
    keywords._p.get_or_add_pPr().append(title_sect)

    _p(doc, styles, "Heading1", "Introduction", 12, "left", True)
    _p(
        doc, styles, "Normal",
        "The November 2023 guidance of the U.S. Food and Drug Administration ties the credibility of a physics-based simulation to the decision it informs, and it places a standalone machine-learning model outside that scope [1]. The American Society of Mechanical Engineers (ASME) V&V 40 asks for the question of interest, the context of use, and the consequence of the decision [2]. NASA-STD-7009B separates the capability of a model from permission to transfer a result [3]. None of these documents returns a grade for one research file. Pathmanathan and colleagues carry the credibility chain down to the submodels of an in silico clinical trial [4]. This paper asks the same question of one auditor: with only the evidence inside the bundle, where does the claim stop?",
        10,
    )
    _p(
        doc, styles, "Normal",
        "Three measurements are reported. A file label is not the audit grade, and a declared model risk does not move that grade. A phenomenon closes only from an executed treatment arm. Six requirement deletions fall below grade 4. No dual extraction of published papers was performed, and no literature rate is stated. The auditor is not an external rater.",
        10,
    )
    _p(
        doc, styles, "Normal",
        "The auditor, ClaimGate, is part of CodonTrace Genesis. The engine does not know a domain, and it was not rewritten for medicine. A domain profile names the port: artificial life, biomedical, or hardware. Records from Avida, a table from MABE2, and an Overview, Design concepts, and Details description are read through that port and are not the product of this paper. An ESP32 port is implemented and is not one of the measurements below. A configuration string that would claim an ASME V&V 40 pass is rejected. This tool does not issue an ASME, FDA, or IEC certificate. Viceconti and colleagues set an in silico trial on context, risk, and a verification chain [5]. Here that chain becomes a grade on one file, not a certificate for a product.",
        10,
    )
    _flow_png(FIG)
    _figure(doc, FIG)
    _p(
        doc, styles, "Caption",
        "Fig. 1. Path from the profile to the grade, and the two recorded grades. The risk bar does not move the grade.",
        8, "center", True,
    )
    _p(
        doc, styles, "Normal",
        "Figure 1 is the path for every measurement below. The lower panel repeats Table 3 as a chart: the executed campaign is grade 4 and the device score table is grade 0. A bar at zero is the result, not a missing drawing. The hardware profile is on the path because the port exists. It supplies none of these grades. The double box is not another grade.",
        10,
    )
    _p(
        doc, styles, "Normal",
        "MIRROR assembles a local evidence bundle for inspection before release and does not offer a regulatory guarantee [6]. NovaFabric records an agent run in a signed capsule [7]. The difference here is a grade computed from the conditions present in the bundle.",
        10,
    )

    _p(doc, styles, "Heading1", "Audit ladder", 12, "left", True)
    _p(
        doc, styles, "Normal",
        "The auditor reads the bundle. It does not call the engine. Each grade requires every condition of the grades below it. Grade 0 means that software identity, source digest, artifact, seed list, configuration fingerprint, artifact manifest, and run record are not all present. Grade 1 still lacks a paired treatment and control, a paired comparison, and a consistent measured difference. Grade 2 still lacks an intervention or a deletion, a negative control, verified replay, a direction of effect, and a written limitation with the artifact. Grade 3 still lacks a numeric interval or at least 16 seeds. Grade 4 has the interval and the seed count. Grade 5 also requires an archive identifier separate from the software bundle, a written limitation, and statistical evidence tied to a deletion. In this implementation the public name of grade 4 is replicated effect. Table 1 names the condition whose absence holds the grade in place.",
        10,
    )
    _p(doc, styles, "Caption", "Table 1. Condition that holds the grade in place", 8, "center", True)
    _table(doc, [
        ["Grade", "If this is missing"],
        ["0", "Version, seeds, fingerprint, manifest, run record"],
        ["1", "Paired arms, paired comparison, consistent difference"],
        ["2", "Intervention, negative control, replay, written limit"],
        ["3", "Numeric interval, or at least 16 seeds"],
        ["4", "Archive identifier separate from the software bundle"],
    ], [1000, 3600])
    _p(
        doc, styles, "Normal",
        "On a compared arm, more outcome rows than seeds void both the interval and the condition of 16 seeds. A subsample is not a new seed. The hash is taken from the file bytes and does not rewrite them. The default test is a sign-flip permutation with a Holm correction and a bias-corrected and accelerated interval, written BCa below. A studentized bootstrap is optional when there are at most 40 pairs.",
        10,
    )

    _p(doc, styles, "Heading1", "Labels are not grades", 12, "left", True)
    _p(
        doc, styles, "Normal",
        "Regulatory words on this port are labels, not grades. ASME VVUQ 40.1-2026 is a worked tibial-tray example, not a pass bit [8]. IEC 62304 assigns software safety class A, B, or C [9]. The International Medical Device Regulators Forum document N12 sets categories I to IV for software as a medical device [10]. Recording a category does not classify a product. If the model is declared not physics-based, the port records that the 2023 guidance places a standalone model outside that scope. The note does not raise the ladder.",
        10,
    )
    _p(
        doc, styles, "Normal",
        "Importance and knowledge are each scored on three levels, following Phenomena Identification and Ranking Table practice [11]. Table 2 is the gap rule on this port. It is not a result from the spent-fuel study in that report. High importance stays open unless knowledge is adequate and the row was measured. Medium importance is a gap only when knowledge is none. Low importance is screened out. A typed word, with no treatment arm, is not a measurement. On this port a coupled model does not rise above its weakest executed submodel. A submodel that is not identifiable cannot contribute above grade 1. Evidence limited to categories 2, 6, and 7 of the 2023 guidance, calibration, emergent behavior, and plausibility, cannot contribute above grade 2 [1]. If a submodel was not executed, that cap is not reported as a measurement and does not move the claim grade.",
        10,
    )
    _p(doc, styles, "Caption", "Table 2. Ranking gap. A typed word is not a measurement", 8, "center", True)
    _table(doc, [
        ["Importance", "Knowledge", "Measured", "Gap"],
        ["High", "None or partial", "Either", "Open"],
        ["High", "Adequate", "No", "Open"],
        ["High", "Adequate", "Yes", "Closes with treatment"],
        ["Medium", "None", "Either", "Open"],
        ["Medium", "Partial or adequate", "Either", "Screened"],
        ["Low", "Any", "Either", "Screened"],
    ], [1300, 1200, 1100, 1000])

    _p(doc, styles, "Heading1", "Two controls", 12, "left", True)
    _p(
        doc, styles, "Normal",
        "In the results file, the label intervention_supported maps to internal grade 3. That name alone does not confer public grade 4. These thresholds belong to this implementation. They are not approval of a device.",
        10,
    )
    _p(
        doc, styles, "Normal",
        "The positive control is the source-bias campaign HE01, replayable, with 30 paired seeds. The results file, fingerprint prefix 28f812c5, contains three comparisons. Against the open channel with no source bias, the difference is 16.49, the raw Monte Carlo p value is 4.99975 × 10⁻⁵, the Holm p value is 1.50 × 10⁻⁴, and the BCa interval is 14.25 to 19.14. The p value is a floor from 20,000 draws, not zero. The primary comparison changes two settings together: the source-fitness threshold moves from 1.5 to 0, and the acceptance policy moves from fitness-weighted to threshold. The difference 16.49 belongs to that pair of settings, not to either setting alone. Against the closed channel and against empty content the difference is 30.67 and the interval is 26.03 to 35.32. In those two contrasts, both content and bias change. The auditor returns grade 4 because replay, the interval, and the seed count hold.",
        10,
    )
    _p(
        doc, styles, "Normal",
        "The negative control is a declared score table for a device model. The question was whether a size choice would be licensed. There is no ISO 14879-1 test and no implant. Model influence is 2 and decision consequence is 3. Treatment scores are 0.12, 0.11, and 0.13. Control scores are 0.20, 0.19, and 0.21. The declared labels are software in a medical device, class B of IEC 62304 [9], category II of IMDRF [10], and evidence categories 1, 3, and 8 of the 2023 guidance, with the model declared as physics-based. Replay is not verified, and the artifact manifest is among the conditions for leaving grade 0, so the grade is 0. That tibial-tray example is not run here [8]. Table 3 records the distance between the two controls.",
        10,
    )
    _p(doc, styles, "Caption", "Table 3. Positive and negative controls", 8, "center", True)
    _table(doc, [
        ["Bundle", "Declared", "Executed", "Grade"],
        ["HE01", "File label", "Replay, 30 seeds, interval", "4"],
        ["Device table", "Regulatory labels", "No recorded run", "0"],
    ], [1000, 1200, 1600, 800])
    _p(
        doc, styles, "Normal",
        "Grade 5 on HE01 lacks only an archived artifact or a DOI for the campaign itself. The software identifier 10.5281/zenodo.20337435 is not counted as that archive, so the positive control stays at grade 4. On these two bundles, changing only declared metadata did not raise the grade. The fields tried were model influence, decision consequence, the software-in-a-device flag, the IEC and IMDRF labels, the FDA evidence categories, and the typed word adequate. Declaring that a model is not physics-based records the scope note in the 2023 guidance and still does not move the ladder.",
        10,
    )

    _p(doc, styles, "Heading1", "Closure and the risk bar", 12, "left", True)
    _p(
        doc, styles, "Normal",
        "The credibility worksheet ranks a phenomenon. The rank alone does not close it. Closure requires the arm to be in the bundle, the role to be treatment, and the arm to be the treatment side of a comparison whose interval bounds are both finite. The auditor must be at grade 4 or higher, and replay must be verified. If the phenomenon names a maximum interval width, the interval must be no wider. One arm closes one phenomenon. A second phenomenon on the same arm is marked shared and stays open. A rank with no arm stays declared.",
        10,
    )
    _p(
        doc, styles, "Normal",
        "On the study bound to HE01, source-fitness bias and the life-loop campaign both close. The life loop is a campaign, not a device, and it carries no FDA evidence category. Contact stress and patient geometry stay declared. Patient geometry remains at a usable level of 2. Because not every submodel was executed, the coupled ceiling is empty. The claim grade of the campaign is still 4. The primary interval is about 4.89 wide. A tolerance of 1.0 does not close source bias. A tolerance of 6.0 does. A closed channel is not a treatment arm. Two submodels that share a configuration fingerprint are not independent and do not form a ceiling. In a separate check, two replayed bundles at grade 4 with different fingerprints can set a ceiling of 4, and that ceiling does not change the public grade. The gap rule is Table 2.",
        10,
    )
    _p(doc, styles, "Caption", "Table 4. Phenomenon closure on the HE01 study", 8, "center", True)
    _table(doc, [
        ["Row", "Evidence", "Result"],
        ["Source bias", "Treatment, grade 4, replay", "Closed"],
        ["Life loop", "Executed campaign", "Closed"],
        ["Contact stress", "Rank, no arm", "Declared only"],
        ["Patient geometry", "Declared submodel, level 2", "Declared only"],
    ], [1300, 2100, 1200])
    _p(
        doc, styles, "Normal",
        "The unit of analysis is the seed, not a row of subsamples. If a compared arm has more values than seeds, the interval and the 16-seed condition both fail, and a pseudoreplication warning is stored [12]. HE01 has one value per seed and remains at grade 4.",
        10,
    )
    _p(
        doc, styles, "Normal",
        "Risk on this port is the maximum of declared influence and declared consequence. The question is the one ASME V&V 40 asks, whether the evidence is commensurate with model risk [2]. The numbers are this implementation, not an FDA table. Risk 1 asks for grade 2, risk 2 for grade 3, and risk 3 for grade 4. From risk 2 upward, an open high-importance phenomenon also blocks the bar. The bar does not call the auditor and does not change the grade. The device table is risk 3 at grade 0, so the bar is not met. HE01 with only the executed source-bias phenomenon meets risk 3 and stays at grade 4. The same campaign, with contact stress added and no arm for it, is still grade 4, and the bar is not met. Declaring the risk neither closes the open phenomenon nor lowers the grade. The committed result has fingerprint prefix b053e7f1. Table 5 is that comparison.",
        10,
    )
    _p(doc, styles, "Caption", "Table 5. Declared risk against the audit grade", 8, "center", True)
    _table(doc, [
        ["Bundle", "Risk", "Grade", "Bar"],
        ["Device table", "3", "0", "Not met"],
        ["HE01, source bias only", "3", "4", "Met"],
        ["HE01 plus contact stress", "3", "4", "Not met"],
    ], [1900, 700, 800, 1200])

    _p(doc, styles, "Heading1", "Requirement deletions", 12, "left", True)
    _p(
        doc, styles, "Normal",
        "The intact reference is grade 4. Six bundles were built by removing one condition, and all six stayed below 4. Identical arms, a zero-width interval, and an unpaired comparison stay at grade 1. Missing replay, and an empty limitation, stay at grade 2. Three seeds stay at grade 3. The six bundles are not an independent random sample, so no binomial interval and no false-accept rate are reported. The count is not a percentage of papers. Table 6 lists the grades. A separate row changes the label of a sensitivity arm and still receives grade 4, because a confirmatory negative-control arm is still in the bundle. The row is a probe, not an escape trial.",
        10,
    )
    _p(doc, styles, "Caption", "Table 6. One condition removed from a grade-4 reference", 8, "center", True)
    _table(doc, [
        ["Deletion", "Grade"],
        ["Identical arms", "1"],
        ["Zero-width interval", "1"],
        ["No paired comparison", "1"],
        ["No replay", "2"],
        ["Empty limitation", "2"],
        ["Three seeds instead of 16", "3"],
    ], [3200, 1400])

    _p(doc, styles, "Heading1", "Arms in the positive control", 12, "left", True)
    _p(
        doc, styles, "Normal",
        "The arm role is read from the campaign, and the results file is not rewritten. Treatment is only the arm with source bias on. Closure asks for that role. A dose arm, a negative control, and a channel that is off do not close a phenomenon, even when they sit in the same grade-4 file. Table 7 lists the roles.",
        10,
    )
    _p(doc, styles, "Caption", "Table 7. Arm role in the positive control", 8, "center", True)
    _table(doc, [
        ["Arm", "Role"],
        ["Source bias on", "Treatment"],
        ["Source bias off", "Mechanism removed"],
        ["Closed channel", "Channel off"],
        ["Empty content", "Negative control"],
        ["Shuffled capsules", "Negative control"],
        ["Matched activity", "Negative control"],
        ["Predictive capsule", "Dose"],
    ], [2300, 2300])
    _p(
        doc, styles, "Normal",
        "The difference 16.49 sets treatment against removal of the mechanism, and both settings change together. The difference 30.67 is the same number for the channel-off arm and for empty content, and those two baselines share one value fingerprint. The width of the first interval is 19.135 minus 14.249, about 4.89. An optional studentized interval on the same 30 differences runs from about 14.05 to 19.30 and does not alter the results file. The study file prefix is d176f0ab and the deletion-file prefix is 3c8a8a1f. Both files were produced by running the auditor. The six deletions in Table 6 stop at the rung that owns the removed condition, and the role-change probe was kept out of that count.",
        10,
    )

    _p(doc, styles, "Heading1", "Discussion", 12, "left", True)
    _p(
        doc, styles, "Normal",
        "The measurements are grades of the evidence inside a bundle. The file label does not confer grade 4. A phenomenon without an executed arm does not close. Deletions of ladder conditions stay below grade 4. The risk bar does not move the grade. A percentage of orthopedic or cardiac claims was not measured. That percentage would need an entry protocol, two extractors, and agreement between raters.",
        10,
    )
    _p(
        doc, styles, "Normal",
        "HE01 is the reference campaign of the auditor. It is not a clinical result and not a device test. The intervention changes two settings together and reports the final energy of a receiver in the simulation. The device table is not a bench test. The deletions follow the auditor's own conditions, so a grade below 4 is the ladder's stated expectation and not an error rate for an outside population. The hardware port is outside these measurements. The software version is read from the project configuration and is not written as a constant inside the bundle.",
        10,
    )
    _p(
        doc, styles, "Normal",
        "Configuration rejects strings that would claim FDA clearance, clinical validation, IEC 62304 certification, or an ASME V&V 40 pass. Rejection is not an audit grade. Four observations can be repeated from the committed files. The campaign audit returns grade 4, and grade 5 lacks only a separate archive identifier. The device table returns grade 0. The committed study and the deletion file match the live auditor. On HE01 the risk bar leaves grade 4 unchanged. The study fingerprint prefix is d176f0ab. The three graded files were committed at c8db778, and those digests are unchanged. The package identity in the project configuration is 0.3.0b7, which is not a PyPI release. The published tip is 0.3.0b6. The DOI named above is the software archive, not the campaign archive.",
        10,
    )

    _p(doc, styles, "Heading1", "Conclusion", 12, "left", True)
    _p(
        doc, styles, "Normal",
        "A file label, an IEC 62304 class, an IMDRF phrase, a typed knowledge word, and a declared risk do not raise the audit grade. The campaign file reaches grade 4 because replay and an interval on 30 seeds hold. The table with no recorded run stays at grade 0. The same grade 4 meets the risk-3 bar when the listed phenomenon was executed, and it does not meet that bar when a high-importance phenomenon is only declared. Six requirement deletions stay below grade 4, and the count is not a literature rate.",
        10,
    )

    _p(doc, styles, "Heading1", "Refused strings and recorded files", 12, "left", True)
    _p(
        doc, styles, "Normal",
        "The biomedical profile refuses a string that would be a certificate. Refusal is not an audit grade. Table 8 lists four of those strings. The same profile still accepts a declared label, such as an IEC class or an IMDRF category, and that label does not raise the grade.",
        10,
    )
    _p(doc, styles, "Caption", "Table 8. Strings the configuration refuses", 8, "center", True)
    _table(doc, [
        ["String", "Result"],
        ["fda_cleared", "Refused"],
        ["clinical_validated", "Refused"],
        ["iec_62304_certified", "Refused"],
        ["asme_vv40_passed", "Refused"],
    ], [2800, 1800])
    _p(
        doc, styles, "Normal",
        "A reader can repeat three files. Table 9 gives the fingerprint prefix of each. The auditor writes the file. The three files were committed at c8db778, and those digests are unchanged. Code is at github.com/Parvaz-Jamei/codontrace-genesis. The package identity is 0.3.0b7 and is not a PyPI release. The published tip is 0.3.0b6. The DOI 10.5281/zenodo.20337435 archives the software, not the campaign.",
        10,
    )
    _p(doc, styles, "Caption", "Table 9. Files a reader can check", 8, "center", True)
    _table(doc, [
        ["File", "Prefix", "Record"],
        ["HE01 study", "d176f0ab", "Grade 4, ceiling empty"],
        ["Deletion grid", "3c8a8a1f", "Six deletions below 4"],
        ["Risk bar", "b053e7f1", "Met only if the row ran"],
    ], [1400, 1400, 1800])
    _p(
        doc, styles, "Normal",
        "Hardware is a third profile on the same port, not a second engine. The ESP32 adapter is an engineering stub. It is implemented, and it is not one of the measurements above. No robot result is reported. A later measurement would use this auditor on that port.",
        10,
    )

    _p(doc, styles, "Heading", "Acknowledgment", 12, "left", True)
    _p(
        doc, styles, "Normal",
        "The biomedical labels follow the FDA, ASME, IEC, and IMDRF documents cited above. The author is solely responsible for the software and for the measurements in this paper. The study grades evidence bundles and does not report a clinical study or a regulatory submission.",
        10,
    )
    _p(doc, styles, "Heading", "References", 12, "left", True)
    refs = [
        "U.S. Food and Drug Administration, “Assessing the credibility of computational modeling and simulation in medical device submissions,” Guidance for Industry and FDA Staff, Nov. 17, 2023.",
        "Assessing Credibility of Computational Modeling through Verification and Validation: Application to Medical Devices, ASME V&V 40-2018. New York: ASME, 2018.",
        "NASA-STD-7009B, Standard for Models and Simulations. Washington, DC: NASA, approved Mar. 5, 2024.",
        "P. Pathmanathan, K. Aycock, A. Badal, R. Bighamian, J. Bodner, B. A. Craven, and S. Niederer, “Credibility assessment of in silico clinical trials for medical devices,” PLoS Comput. Biol., vol. 20, no. 8, Art. no. e1012289, Aug. 2024, doi: 10.1371/journal.pcbi.1012289.",
        "M. Viceconti, F. Pappalardo, B. Rodriguez, M. Horner, J. Bischoff, and F. Musuamba Tshinanu, “In silico trials: verification, validation and uncertainty quantification of predictive models used in the regulatory evaluation of biomedical products,” Methods, vol. 185, pp. 120-127, Jan. 2021, doi: 10.1016/j.ymeth.2020.01.011.",
        "Anton Sokolov, “MIRROR: local evidence bundles for reproducible research software review,” Zenodo, May 30, 2026, doi: 10.5281/zenodo.20463358.",
        "M. Seyedkazemi Ardebili, “NovaFabric: tamper-evident, replayable evidence for autonomous AI agent runs,” arXiv:2609.12582, Sep. 11, 2026, doi: 10.48550/arXiv.2609.12582.",
        "An Example of Assessing Computational Model Credibility Using the ASME V&V 40 Risk-Based Framework: Tibial Tray Component Worst-Case Size Identification for Fatigue Testing, ASME VVUQ 40.1-2026. New York: ASME, 2026.",
        "IEC 62304:2006+AMD1:2015, Medical device software, Software life cycle processes. Geneva: IEC, 2015.",
        "IMDRF/SaMD WG/N12FINAL:2014, “Software as a Medical Device”: Possible Framework for Risk Categorization and Corresponding Considerations. IMDRF, 2014.",
        "Nuclear Energy Agency, Phenomena Identification and Ranking Table: R&D Priorities for Loss-of-Cooling and Loss-of-Coolant Accidents in Spent Nuclear Fuel Pools, NEA/CSNI/R(2017)18. Paris: OECD, 2018.",
        "S. H. Hurlbert, “Pseudoreplication and the design of ecological field experiments,” Ecological Monographs, vol. 54, no. 2, pp. 187-211, 1984, doi: 10.2307/1942661.",
    ]
    for ref in refs:
        _p(doc, styles, "References", ref, 8, "both")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    _sanitize(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
